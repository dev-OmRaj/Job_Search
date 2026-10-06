from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .models import ProjectPreferences, ResumeProfile
from .progress import ProgressCallback, emit_progress

SUPPORTED_RESUME_TYPES = {".pdf", ".docx", ".txt"}

EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
PHONE_RE = re.compile(r"(?<!\d)(?:\+?\d[\d\s().-]{7,}\d)(?!\d)")
URL_RE = re.compile(r"\b(?:https?://|www\.)\S+", re.IGNORECASE)
PII_LABEL_RE = re.compile(
    r"^\s*(?:name|address|phone|mobile|email|e-mail|linkedin|github|portfolio)\s*[:|-]",
    re.IGNORECASE,
)


@dataclass(slots=True)
class ResumeLoadReport:
    profiles: list[ResumeProfile]
    errors: list[str] = field(default_factory=list)


def _read_pdf(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def _read_docx(path: Path) -> str:
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs)


def read_resume(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_RESUME_TYPES:
        supported = ", ".join(sorted(SUPPORTED_RESUME_TYPES))
        raise ValueError(f"Unsupported resume type {suffix!r}; supported: {supported}")
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    return path.read_text(encoding="utf-8")


def sanitize_resume_text(text: str, *, max_characters: int = 12_000) -> str:
    """Best-effort PII removal before any resume-derived text reaches the cloud."""

    retained_lines: list[str] = []
    for line in text.splitlines():
        if PII_LABEL_RE.search(line):
            continue
        line = EMAIL_RE.sub("[email removed]", line)
        line = PHONE_RE.sub("[phone removed]", line)
        line = URL_RE.sub("[url removed]", line)
        normalized = " ".join(line.split())
        if normalized:
            retained_lines.append(normalized)
    return "\n".join(retained_lines)[:max_characters]


def _load_profile(name: str, path: Path, preferences) -> ResumeProfile:
    raw_bytes = path.read_bytes()
    raw_text = read_resume(path)
    summary = sanitize_resume_text(raw_text)
    if not summary:
        raise ValueError(f"No readable resume text found for profile {name!r}: {path}")
    return ResumeProfile(
        name=name,
        path=path,
        file_hash=hashlib.sha256(raw_bytes).hexdigest(),
        sanitized_summary=summary,
        preferences=preferences,
    )


def _discovered_resume_paths(resume_dir: Path) -> list[Path]:
    if not resume_dir.is_dir():
        return []
    return sorted(
        (
            path
            for path in resume_dir.iterdir()
            if path.is_file() and path.suffix.lower() in SUPPORTED_RESUME_TYPES
        ),
        key=lambda path: (path.name.casefold(), path.name),
    )


def _profile_names(paths: list[Path]) -> dict[Path, str]:
    stem_counts: dict[str, int] = {}
    for path in paths:
        key = path.stem.casefold()
        stem_counts[key] = stem_counts.get(key, 0) + 1

    names: dict[Path, str] = {}
    used: set[str] = set()
    for path in paths:
        name = path.stem
        if stem_counts[path.stem.casefold()] > 1:
            name = f"{path.stem}_{path.suffix.lower().lstrip('.')}"
        candidate = name
        sequence = 2
        while candidate.casefold() in used:
            candidate = f"{name}_{sequence}"
            sequence += 1
        names[path] = candidate
        used.add(candidate.casefold())
    return names


def load_resume_profiles(
    preferences: ProjectPreferences,
    resume_dir: Path,
    progress: ProgressCallback | None = None,
) -> list[ResumeProfile]:
    """Strictly load every supported resume discovered directly under the folder."""

    report = load_active_resume_profiles(preferences, resume_dir, progress)
    if report.errors:
        raise ValueError("; ".join(report.errors))
    if not report.profiles:
        supported = ", ".join(sorted(SUPPORTED_RESUME_TYPES))
        raise FileNotFoundError(
            f"No supported resumes found in {resume_dir.resolve()}; add one of: {supported}"
        )
    return report.profiles


def load_active_resume_profiles(
    preferences: ProjectPreferences,
    resume_dir: Path,
    progress: ProgressCallback | None = None,
) -> ResumeLoadReport:
    """Discover and load every usable resume directly under ``resume_dir``."""

    root = resume_dir.resolve()
    profiles: list[ResumeProfile] = []
    errors: list[str] = []
    paths = _discovered_resume_paths(resume_dir)
    profile_names = _profile_names(paths)
    emit_progress(
        progress,
        "resumes",
        "Loading resumes",
        completed=0,
        total=len(paths),
    )
    for index, discovered_path in enumerate(paths, start=1):
        path = discovered_path.resolve()
        name = profile_names[discovered_path]
        if not path.is_relative_to(root):
            errors.append(f"Ignored resume {discovered_path.name!r}: it resolves outside {root}")
            emit_progress(
                progress,
                "resumes",
                "Loading resumes",
                completed=index,
                total=len(paths),
                finished=index == len(paths),
            )
            continue
        try:
            profiles.append(_load_profile(name, path, preferences.job_preferences))
        except Exception as exc:
            errors.append(f"Ignored resume {discovered_path.name!r}: {exc}")
        emit_progress(
            progress,
            "resumes",
            "Loading resumes",
            completed=index,
            total=len(paths),
            finished=index == len(paths),
        )

    if not paths:
        emit_progress(
            progress,
            "resumes",
            "No resumes found",
            completed=0,
            total=0,
            finished=True,
        )

    return ResumeLoadReport(profiles=profiles, errors=errors)
